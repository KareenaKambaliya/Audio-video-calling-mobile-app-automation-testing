"""
Excel Reader Module.
Reads and parses test cases from test data Excel workbooks across all sheets for data-driven testing.
"""

from pathlib import Path
import openpyxl

from config import config
from utils.logger import logger

DATA_DIR = config.ROOT_DIR / "data"
DEFAULT_TEST_DATA_FILE = getattr(config, "TEST_DATA_FILE", DATA_DIR / "testcase_data.xlsx")
if not DEFAULT_TEST_DATA_FILE.exists():
    DEFAULT_TEST_DATA_FILE = DATA_DIR / "login_test_data.xlsx"


def get_sheet_test_cases(sheet_name: str, file_path: Path = None, only_enabled: bool = False) -> list[dict]:
    """
    Read test scenarios from a specific sheet in the Excel workbook.

    :param sheet_name: Name of the sheet (e.g. 'Login_Test_Cases', 'Create_Call_Test_Cases').
    :param file_path: Path to Excel file (defaults to data/testcase_data.xlsx).
    :param only_enabled: If True, only returns rows where 'Execute' is 'Y'.
    :return: List of dictionaries with test case fields.
    """
    target_path = Path(file_path or DEFAULT_TEST_DATA_FILE)

    if not target_path.exists():
        logger.error(f"Test data file not found: {target_path}")
        return []

    wb = openpyxl.load_workbook(target_path, data_only=True)
    if sheet_name not in wb.sheetnames:
        # Fallback to active sheet if sheet_name not found
        logger.warning(f"Sheet '{sheet_name}' not found in {target_path.name}. Available: {wb.sheetnames}. Using active sheet.")
        ws = wb.active
    else:
        ws = wb[sheet_name]

    headers = [str(cell.value or "").strip() for cell in ws[1]]
    while headers and not headers[-1]:
        headers.pop()

    test_cases = []
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not any(row):
            continue  # Skip completely empty rows

        row_dict = {}
        for header, val in zip(headers, row):
            row_dict[header] = "" if val is None else str(val).strip()

        # Check execution flag
        execute_flag = row_dict.get("Execute", "Y").upper()
        if only_enabled and execute_flag != "Y":
            continue

        test_cases.append(row_dict)

    logger.info(f"Loaded {len(test_cases)} test case(s) from [{sheet_name}] in {target_path.name}")
    return test_cases


def get_login_test_cases(file_path: Path = None, only_enabled: bool = False) -> list[dict]:
    """Backward-compatible helper for login test scenarios."""
    cases = get_sheet_test_cases("Login_Test_Cases", file_path=file_path, only_enabled=only_enabled)
    if not cases:
        # Fallback to default active sheet (e.g. login_test_data.xlsx)
        cases = get_sheet_test_cases(None, file_path=DATA_DIR / "login_test_data.xlsx", only_enabled=only_enabled)
    return cases


def get_primary_credentials(file_path: Path = None) -> tuple[str, str]:
    """
    Get the first active test case credentials from the Excel sheet.
    Falls back to config default if none found.
    """
    cases = get_login_test_cases(file_path=file_path, only_enabled=True)
    if cases:
        first = cases[0]
        return first.get("Email", config.LOGIN_EMAIL), first.get("Password", config.LOGIN_PASSWORD)
    return config.LOGIN_EMAIL, config.LOGIN_PASSWORD
